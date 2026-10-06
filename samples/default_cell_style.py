"""Generate a maintained sample preserving default fills for absent blank cells."""

from __future__ import annotations

from copy import copy
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.indexed_list import IndexedList

from excel_template_writer.xlsx import render_workbook
from samples._common import (
    GRID_BORDER,
    LIGHT_BLUE,
    LIGHT_GOLD,
    LIGHT_GREEN,
    NAVY,
    SAMPLES_DIR,
    WHITE,
    assert_no_template_tags,
    atomic_save,
    paint,
    sample_paths,
)

TEMPLATE_PATH, OUTPUT_PATH = sample_paths("default_cell_style")
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
TEXT_ROWS = {6: "Text 1 - overview", 8: "Text 2 - details", 10: "Text 3 - next steps"}
TABLES = (
    (14, "table_one", (("Apples", 3), ("Pears", 2), ("Oranges", 5)), LIGHT_BLUE),
    (19, "table_two", (("Notebook", 4), ("Pen", 12)), LIGHT_GREEN),
    (24, "table_three", (("Alpha", 1), ("Beta", 2), ("Gamma", 3), ("Delta", 4)), LIGHT_GOLD),
)


def build_template(path: Path = TEMPLATE_PATH) -> Path:
    """Build visible table tags beneath an implicitly white blank band.

    Args:
        path: Destination path for the authored template workbook.

    Returns:
        The verified template path.
    """

    workbook = Workbook()
    white_fill = PatternFill("solid", fgColor=WHITE)
    # Excel can encode a workbook's white background in cellXfs[0], with no
    # individual cell records for its blank areas. Reindex after changing the
    # default so openpyxl's hashed style table remains internally consistent.
    workbook._cell_styles[0].fillId = workbook._fills.add(copy(white_fill))
    workbook._cell_styles = IndexedList([copy(workbook._cell_styles[0])])
    workbook._named_styles[0].fill = copy(white_fill)
    sheet = workbook.active
    sheet.title = "Default cell style"
    sheet.sheet_view.showGridLines = True
    sheet.sheet_view.zoomScale = 80
    sheet.sheet_format.defaultRowHeight = 23
    for column, width in zip("ABCDEFGH", (40, 27, 23, 22, 10, 10, 10, 10), strict=True):
        sheet.column_dimensions[column].width = width
    sheet["A1"] = "Workbook default cell style"
    sheet["A1"].font = Font(name="Aptos", size=18, bold=True, color=NAVY)
    sheet["A1"].fill = copy(white_fill)
    sheet.row_dimensions[1].height = 32
    sheet["A2"] = "B6:H11 stays white with gridlines on, without individual blank-cell records."
    sheet["A3"] = "The workbook default supplies the white fill; the three tables expand below."
    sheet["A4"] = "Row 12 controls: B:D direct white fill; E:G explicit no fill, all with borders."
    for row in (2, 3, 4):
        sheet.cell(row, 1).fill = copy(white_fill)
        sheet.cell(row, 1).font = Font(name="Aptos", size=11, color=NAVY)
    # Do not materialize B6:H11: their effective fill must come from the workbook
    # default, and cannot be tested by reading openpyxl's synthesized cells.
    for row, text in TEXT_ROWS.items():
        sheet.cell(row, 1, text)
        sheet.cell(row, 1).fill = copy(white_fill)
        sheet.cell(row, 1).font = Font(name="Aptos", size=12, color=NAVY)
    sheet["A12"] = "Direct fill / no-fill controls"
    sheet["A12"].fill = copy(white_fill)
    paint(sheet, "B12:D12", fill=WHITE)
    paint(sheet, "E12:G12", fill=WHITE)
    for column in "EFG":
        sheet[f"{column}12"].fill = PatternFill()
    for index, (header_row, key, _, fill) in enumerate(TABLES, start=1):
        labels = (f"Table {index} / start tag", "Item / expression", "Qty / expression", "End tag")
        for column, label in enumerate(labels, start=1):
            sheet.cell(header_row, column, label)
        paint(sheet, f"A{header_row}:D{header_row}", fill=NAVY, bold=True, font_color=WHITE)
        tags = (f"{{% for item in {key} %}}", "{{ item.name }}", "{{ item.qty }}", "{% endfor %}")
        for column, tag in enumerate(tags, start=1):
            sheet.cell(header_row + 1, column, tag)
        paint(sheet, f"A{header_row + 1}:D{header_row + 1}", fill=fill)
        sheet.row_dimensions[header_row].height = 25
        sheet.row_dimensions[header_row + 1].height = 25
    # Width/height dimensions must not accidentally override the customized
    # workbook default with openpyxl's ordinary all-zero, no-fill cell style.
    for dimension in (*sheet.row_dimensions.values(), *sheet.column_dimensions.values()):
        dimension._style = copy(workbook._cell_styles[0])
    result = atomic_save(workbook, path)
    workbook.close()
    _assert_implicit_white_band(path)
    return result


def _assert_implicit_white_band(path: Path) -> None:
    """Verify both workbook defaults and the absence of blank-cell records.

    Args:
        path: Template or rendered workbook to inspect without materializing cells.
    """

    with ZipFile(path) as archive:
        styles = ET.fromstring(archive.read("xl/styles.xml"))
        sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    fills = styles.find("s:fills", NS)
    cell_styles = styles.find("s:cellXfs", NS)
    named_styles = styles.find("s:cellStyles", NS)
    named_xfs = styles.find("s:cellStyleXfs", NS)
    assert fills is not None and cell_styles is not None
    assert named_styles is not None and named_xfs is not None
    normal_ref = next(style for style in named_styles if style.get("builtinId") == "0")
    normal_xf_id = normal_ref.get("xfId")
    assert normal_xf_id is not None
    for xf in (cell_styles[0], named_xfs[int(normal_xf_id)]):
        default_fill = fills[int(xf.get("fillId", "0"))].find("s:patternFill", NS)
        assert default_fill is not None
        assert default_fill.get("patternType") == "solid"
        foreground = default_fill.find("s:fgColor", NS)
        assert foreground is not None and foreground.get("rgb") == WHITE
    stored_cells = {cell.attrib["r"] for cell in sheet.findall("s:sheetData/s:row/s:c", NS)}
    blank_band = {f"{column}{row}" for row in range(6, 12) for column in "BCDEFGH"}
    assert blank_band.isdisjoint(stored_cells)
    # Row/column overrides would change the effective appearance of absent cells.
    assert all("s" not in row.attrib for row in sheet.findall("s:sheetData/s:row", NS))
    assert all("style" not in column.attrib for column in sheet.findall("s:cols/s:col", NS))


def render_sample(
    template_path: Path = TEMPLATE_PATH,
    output_path: Path = OUTPUT_PATH,
) -> Path:
    """Render tables and verify implicit defaults alongside direct presentation.

    Args:
        template_path: Authored sample template path.
        output_path: Separate rendered workbook path.

    Returns:
        The verified rendered workbook path.
    """

    context = {
        key: [{"name": name, "qty": quantity} for name, quantity in items]
        for _, key, items, _ in TABLES
    }
    render_workbook(template_path, output_path, context)
    _assert_implicit_white_band(template_path)
    _assert_implicit_white_band(output_path)
    assert_no_template_tags(output_path)
    template, output = load_workbook(template_path), load_workbook(output_path)
    try:
        assert template.sheetnames == output.sheetnames == ["Default cell style"]
        source, result = template.active, output.active
        assert source.sheet_view.showGridLines is result.sheet_view.showGridLines is True
        assert source.sheet_view.zoomScale == result.sheet_view.zoomScale == 80
        assert not source.merged_cells and not result.merged_cells
        for row, text in TEXT_ROWS.items():
            assert source.cell(row, 1).value == result.cell(row, 1).value == text
            assert result.cell(row, 1).fill.fill_type == "solid"
            assert result.cell(row, 1).fill.fgColor.rgb == WHITE
        for sheet in (source, result):
            for column in "BCD":
                assert sheet[f"{column}12"].fill.fill_type == "solid"
                assert sheet[f"{column}12"].fill.fgColor.rgb == WHITE
            for column in "EFG":
                assert sheet[f"{column}12"].fill.fill_type is None
            for column in "BCDEFG":
                assert sheet[f"{column}12"].value is None
                assert sheet[f"{column}12"].border == GRID_BORDER
        shift = 0
        for header_row, key, items, fill in TABLES:
            assert source.cell(header_row + 1, 1).value == f"{{% for item in {key} %}}"
            assert source.cell(header_row + 1, 4).value == "{% endfor %}"
            output_start = header_row + shift + 1
            for offset, (name, quantity) in enumerate(items):
                row = output_start + offset
                assert result.cell(row, 1).value is None
                assert result.cell(row, 2).value == name
                assert result.cell(row, 3).value == quantity
                assert result.cell(row, 3).data_type == "n"
                assert result.cell(row, 4).value is None
                assert result.row_dimensions[row].height == 25
                for column in range(1, 5):
                    cell = result.cell(row, column)
                    assert cell.fill.fill_type == "solid"
                    assert cell.fill.fgColor.rgb == fill
                    assert cell.border == GRID_BORDER
            shift += len(items) - 1
    finally:
        template.close()
        output.close()
    return output_path


def main() -> None:
    """Generate the default-style sample template and its verified output."""

    template = build_template()
    output = render_sample(template)
    print(f"Template: {template.relative_to(SAMPLES_DIR.parent)}")
    print(f"Output:   {output.relative_to(SAMPLES_DIR.parent)}")


if __name__ == "__main__":
    main()
