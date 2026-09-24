"""Generate the maintained sample for fixed no-shift destination filling."""

from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection

from excel_template_writer.diagnostics import DiagnosticCode, TemplateRenderError
from excel_template_writer.xlsx import render_workbook
from samples._common import (
    LIGHT_BLUE,
    LIGHT_GOLD,
    LIGHT_GRAY,
    NAVY,
    SAMPLES_DIR,
    WHITE,
    assert_no_template_tags,
    atomic_save,
    paint,
    prepare_sheet,
    sample_paths,
)

TEMPLATE_PATH, OUTPUT_PATH = sample_paths("fixed_layout_fill")


def build_template(path: Path = TEMPLATE_PATH) -> Path:
    """Build an alternating fixed-row template.

    Args:
        path: Destination path for the authored template.

    Returns:
        The verified template path.
    """

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fixed layout fill"
    prepare_sheet(
        sheet,
        "Fixed no-shift destination fill",
        "Values fill prepared alternating rows; their colors, formats, and heights stay authored.",
        widths=(25, 18, 16, 24),
    )
    for column, label in enumerate(("Item", "Amount", "State", "Stationary label"), start=1):
        sheet.cell(4, column, label)
    paint(sheet, "A4:D4", fill=NAVY, bold=True, font_color=WHITE, horizontal="center")

    sheet["A5"] = '{% for item in items shift="none" %}{{ item.label }}'
    sheet["B5"] = "{{ item.amount }}"
    sheet["C5"] = "{{ item.state }}{% endfor %}"
    fills = (LIGHT_BLUE, WHITE, LIGHT_BLUE, WHITE, LIGHT_BLUE, WHITE)
    heights = (22, 25, 23, 26, 24, 27)
    for row, (fill, height) in enumerate(zip(fills, heights, strict=True), start=5):
        paint(sheet, f"A{row}:C{row}", fill=fill, bold=fill == LIGHT_BLUE)
        sheet.row_dimensions[row].height = height
        sheet[f"B{row}"].number_format = (
            "$#,##0.00;[Red]-$#,##0.00" if fill == LIGHT_BLUE else "0.000"
        )
        for column in range(1, 4):
            sheet.cell(row, column).protection = Protection(locked=fill == LIGHT_BLUE)
        sheet[f"D{row}"] = f"Stays on row {row}"
        paint(sheet, f"D{row}", fill=LIGHT_GRAY)

    sheet["A11"] = "Prepared-band footer — stays on row 11"
    sheet["B11"] = "Overflow collides here"
    sheet["C11"] = "No movement"
    sheet["D11"] = "Stationary footer"
    paint(sheet, "A11:D11", fill=LIGHT_GOLD, bold=True)
    result = atomic_save(workbook, path)
    workbook.close()
    return result


def _items(count: int) -> list[dict[str, object]]:
    """Return deterministic sample rows in display order.

    Args:
        count: Number of rows to create.

    Returns:
        Prepared records in their intended display order.
    """

    return [
        {
            "label": f"Prepared item {index}",
            "amount": index * 125.5,
            "state": "Open" if index % 2 else "Closed",
        }
        for index in range(1, count + 1)
    ]


def render_sample(
    template_path: Path = TEMPLATE_PATH,
    output_path: Path = OUTPUT_PATH,
) -> Path:
    """Render and verify fixed no-shift placement.

    Args:
        template_path: Authored sample template path.
        output_path: Separate rendered workbook path.

    Returns:
        The verified rendered workbook path.
    """

    render_workbook(template_path, output_path, {"items": _items(4)})
    assert_no_template_tags(output_path)
    workbook = load_workbook(output_path)
    try:
        sheet = workbook["Fixed layout fill"]
        assert [sheet[f"A{row}"].value for row in range(5, 11)] == [
            "Prepared item 1",
            "Prepared item 2",
            "Prepared item 3",
            "Prepared item 4",
            None,
            None,
        ]
        assert [sheet[f"A{row}"].fill.fgColor.rgb for row in range(5, 11)] == [
            LIGHT_BLUE,
            WHITE,
            LIGHT_BLUE,
            WHITE,
            LIGHT_BLUE,
            WHITE,
        ]
        assert [sheet[f"B{row}"].number_format for row in range(5, 9)] == [
            "$#,##0.00;[Red]-$#,##0.00",
            "0.000",
            "$#,##0.00;[Red]-$#,##0.00",
            "0.000",
        ]
        assert [sheet.row_dimensions[row].height for row in range(5, 11)] == [
            22,
            25,
            23,
            26,
            24,
            27,
        ]
        assert sheet["A11"].value.startswith("Prepared-band footer")
        assert sheet["D5"].value == "Stays on row 5"
        assert sheet["D11"].value == "Stationary footer"
    finally:
        workbook.close()

    with TemporaryDirectory(dir=output_path.parent) as directory:
        overflow_path = Path(directory) / "fixed_layout_fill_overflow.xlsx"
        try:
            render_workbook(template_path, overflow_path, {"items": _items(7)})
        except TemplateRenderError as error:
            assert DiagnosticCode.LAYOUT_COLLISION in {
                diagnostic.code for diagnostic in error.diagnostics
            }
        else:
            raise AssertionError("fixed no-shift overflow should collide with the footer")
        assert not overflow_path.exists()
    return output_path


def main() -> None:
    """Generate both fixed-fill sample workbooks and print their paths."""

    template = build_template()
    output = render_sample(template)
    print(f"Template: {template.relative_to(SAMPLES_DIR.parent)}")
    print(f"Output:   {output.relative_to(SAMPLES_DIR.parent)}")


if __name__ == "__main__":
    main()
