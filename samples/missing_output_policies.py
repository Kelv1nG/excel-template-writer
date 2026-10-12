"""Generate maintained samples for missing output and collection policies."""

from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from excel_template_writer.diagnostics import (
    Diagnostic,
    DiagnosticCode,
    DiagnosticSeverity,
    TemplateRenderError,
)
from excel_template_writer.xlsx import render_workbook
from samples._common import (
    LIGHT_BLUE,
    LIGHT_GOLD,
    LIGHT_GREEN,
    SAMPLES_DIR,
    atomic_save,
    paint,
    prepare_sheet,
    sample_paths,
)

TEMPLATE_PATH, OUTPUT_PATH = sample_paths("missing_output_policies")
EXACT_MISSING_TAG = "{{  customer.email  }}"
MIXED_MISSING_VALUE = f"Contact: {EXACT_MISSING_TAG}"
NULL_AGGREGATE_TAG = "{{ null_rows | sum }}"
MISSING_REPEAT_OPEN = '{% for row in tables["some_table"] %}{{ row.name }}'
NULL_REPEAT_OPEN = "{% for row in null_rows %}{{ row.name }}"
REPEAT_AMOUNT_TAG = "{{ row.amount }}"
REPEAT_CLOSE = "{% endfor %}"


def build_template(path: Path = TEMPLATE_PATH) -> Path:
    """Build a template showing strict, blank, and preserve recovery policies.

    Args:
        path: Destination path for the authored template workbook.

    Returns:
        The verified template path.
    """

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Missing outputs"
    prepare_sheet(
        sheet,
        "Missing output policies",
        "Strict mode rejects unavailable values; blank and preserve modes publish warnings.",
        widths=(28, 42, 54, 18),
    )
    for column, label in enumerate(
        ("Case", "Template value", "Expected preserve-mode result"),
        start=1,
    ):
        sheet.cell(4, column, label)
    paint(sheet, "A4:C4", fill=LIGHT_BLUE, bold=True, horizontal="center")
    rows = (
        ("Provided value", "{{ customer.name }}", "Acme Industries"),
        ("Missing sole expression", EXACT_MISSING_TAG, "Exact authored tag, including spaces"),
        ("Missing in mixed text", MIXED_MISSING_VALUE, "Literal prefix plus exact authored tag"),
        (
            "Explicit default",
            '{{ customer.phone | default("Not supplied") }}',
            "Not supplied; no warning",
        ),
        ("Present null", "{{ customer.middle_name }}", "Blank; null is not missing"),
    )
    for row_number, values in enumerate(rows, start=5):
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
        paint(
            sheet,
            f"A{row_number}:C{row_number}",
            fill=LIGHT_GOLD if row_number % 2 else "FFFFFFFF",
        )
    paint(sheet, "B6:B7", fill=LIGHT_GREEN)
    sheet.row_dimensions[7].height = 30
    sheet["A10"] = "Null collection aggregate"
    sheet["B10"] = NULL_AGGREGATE_TAG
    sheet["C10"] = "Exact tag remains visible in preserve mode"
    sheet["A12"] = "Missing table repeat"
    sheet["B12"] = MISSING_REPEAT_OPEN
    sheet.merge_cells("C12:D12")
    sheet["B13"] = "Static body"
    sheet["C13"] = REPEAT_AMOUNT_TAG
    sheet["D13"] = REPEAT_CLOSE
    sheet["A14"] = "Between blocks"
    sheet["D14"] = "Unaffected"
    sheet["A16"] = "Null collection repeat"
    sheet["B16"] = NULL_REPEAT_OPEN
    sheet.merge_cells("C16:D16")
    sheet["B17"] = "Static body"
    sheet["C17"] = REPEAT_AMOUNT_TAG
    sheet["D17"] = REPEAT_CLOSE
    paint(sheet, "B12:D13", fill=LIGHT_GOLD)
    paint(sheet, "B16:D17", fill=LIGHT_GOLD)
    paint(sheet, "B12", fill=LIGHT_GREEN)
    paint(sheet, "B16", fill=LIGHT_GREEN)
    sheet.row_dimensions[13].height = 26
    sheet.row_dimensions[17].height = 26
    result = atomic_save(workbook, path)
    workbook.close()
    return result


def _context() -> dict[str, object]:
    """Return deliberately incomplete input shared by all three policy renders."""

    return {
        "customer": {
            "name": "Acme Industries",
            "middle_name": None,
        },
        "null_rows": None,
    }


def _assert_missing_diagnostics(
    diagnostics: tuple[Diagnostic, ...],
    *,
    tolerant: bool,
) -> None:
    """Assert each unavailable output or repeat has a located diagnostic.

    Args:
        diagnostics: Render diagnostics to verify.
        tolerant: Whether blank or preserve recovery was requested.
    """

    missing_code = (
        DiagnosticCode.MISSING_VALUE_RENDERED if tolerant else DiagnosticCode.MISSING_VALUE
    )
    null_filter_code = (
        DiagnosticCode.NULL_COLLECTION_RENDERED if tolerant else DiagnosticCode.FILTER_TYPE_MISMATCH
    )
    null_repeat_code = (
        DiagnosticCode.NULL_COLLECTION_RENDERED if tolerant else DiagnosticCode.EXPECTED_COLLECTION
    )
    expected = [
        (missing_code, "Missing outputs!B6:0"),
        (missing_code, "Missing outputs!B7:9"),
        (null_filter_code, "Missing outputs!B10:0"),
        (missing_code, "Missing outputs!B12:0"),
        (null_repeat_code, "Missing outputs!B16:0"),
    ]
    assert sorted(
        (diagnostic.code, str(diagnostic.location)) for diagnostic in diagnostics
    ) == sorted(expected)
    severity = DiagnosticSeverity.WARNING if tolerant else DiagnosticSeverity.ERROR
    assert all(diagnostic.severity is severity for diagnostic in diagnostics)
    assert sum("customer.email" in diagnostic.message for diagnostic in diagnostics) == 2


def _assert_common_values(path: Path) -> None:
    """Reopen one rendered workbook and verify non-missing value behavior.

    Args:
        path: Rendered workbook path to reopen.
    """

    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook["Missing outputs"]
        assert sheet["B5"].value == "Acme Industries"
        assert sheet["B8"].value == "Not supplied"
        assert sheet["B9"].value is None
        assert sheet["A14"].value == "Between blocks"
        assert sheet["D14"].value == "Unaffected"
        assert {str(merge) for merge in sheet.merged_cells.ranges} == {
            "A1:D1",
            "A2:D2",
            "C12:D12",
            "C16:D16",
        }
        assert sheet.row_dimensions[13].height == 26
        assert sheet.row_dimensions[17].height == 26
        assert sheet.column_dimensions["D"].width == 18
    finally:
        workbook.close()


def _verify_strict_mode(template_path: Path, output_path: Path) -> None:
    """Verify strict mode reports errors and publishes no workbook.

    Args:
        template_path: Authored workbook containing unavailable values.
        output_path: Temporary path that must remain unpublished.
    """

    try:
        render_workbook(
            template_path,
            output_path,
            _context(),
            missing_output="error",
        )
    except TemplateRenderError as error:
        _assert_missing_diagnostics(error.diagnostics, tolerant=False)
    else:
        raise AssertionError("strict missing-output mode must reject the render")
    assert not output_path.exists()


def _verify_blank_mode(template_path: Path, output_path: Path) -> None:
    """Verify blank mode publishes formatted blanks and warning diagnostics.

    Args:
        template_path: Authored workbook containing unavailable values.
        output_path: Temporary blank-mode output path.
    """

    result = render_workbook(
        template_path,
        output_path,
        _context(),
        missing_output="blank",
    )
    _assert_missing_diagnostics(result.diagnostics, tolerant=True)
    _assert_common_values(output_path)
    workbook = load_workbook(output_path, data_only=False)
    try:
        sheet = workbook["Missing outputs"]
        assert sheet["B6"].value is None
        assert sheet["B6"].fill.fgColor.rgb == LIGHT_GREEN
        assert sheet["B7"].value == "Contact: "
        assert sheet["B10"].value is None
        for coordinate in ("B12", "B13", "C13", "D13", "B16", "B17", "C17", "D17"):
            assert sheet[coordinate].value is None
        assert sheet["B12"].fill.fgColor.rgb == LIGHT_GREEN
        assert sheet["B16"].fill.fgColor.rgb == LIGHT_GREEN
        assert sheet["C13"].fill.fgColor.rgb == LIGHT_GOLD
    finally:
        workbook.close()


def render_sample(
    template_path: Path = TEMPLATE_PATH,
    output_path: Path = OUTPUT_PATH,
) -> Path:
    """Render and verify strict, blank, and preserve policies.

    The committed output workbook uses preserve mode so unresolved output tags
    and repeat rectangles remain directly inspectable. Strict and blank outputs
    are temporary.

    Args:
        template_path: Authored sample template path.
        output_path: Separate preserve-mode output workbook path.

    Returns:
        The verified preserve-mode output path.
    """

    with TemporaryDirectory() as temporary_directory:
        temporary_path = Path(temporary_directory)
        _verify_strict_mode(template_path, temporary_path / "strict.xlsx")
        _verify_blank_mode(template_path, temporary_path / "blank.xlsx")

    result = render_workbook(
        template_path,
        output_path,
        _context(),
        missing_output="preserve",
    )
    _assert_missing_diagnostics(result.diagnostics, tolerant=True)
    _assert_common_values(output_path)
    workbook = load_workbook(output_path, data_only=False)
    try:
        sheet = workbook["Missing outputs"]
        assert sheet["B6"].value == EXACT_MISSING_TAG
        assert sheet["B6"].fill.fgColor.rgb == LIGHT_GREEN
        assert sheet["B7"].value == MIXED_MISSING_VALUE
        assert sheet["B10"].value == NULL_AGGREGATE_TAG
        assert sheet["B12"].value == MISSING_REPEAT_OPEN
        assert sheet["B13"].value == "Static body"
        assert sheet["C13"].value == REPEAT_AMOUNT_TAG
        assert sheet["D13"].value == REPEAT_CLOSE
        assert sheet["B16"].value == NULL_REPEAT_OPEN
        assert sheet["B17"].value == "Static body"
        assert sheet["C17"].value == REPEAT_AMOUNT_TAG
        assert sheet["D17"].value == REPEAT_CLOSE
        assert sheet["B12"].fill.fgColor.rgb == LIGHT_GREEN
        assert sheet["B16"].fill.fgColor.rgb == LIGHT_GREEN
    finally:
        workbook.close()
    return output_path


def main() -> None:
    """Generate the missing-output template and preserve-mode output workbooks."""

    template = build_template()
    output = render_sample(template)
    print(f"Template: {template.relative_to(SAMPLES_DIR.parent)}")
    print(f"Output:   {output.relative_to(SAMPLES_DIR.parent)}")


if __name__ == "__main__":
    main()
