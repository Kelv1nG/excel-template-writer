"""Characterize the pinned hook and isolate process-lifetime activation in children."""

import math
import os
import random
import struct
import subprocess
import sys
import textwrap
from collections.abc import Sequence
from contextlib import nullcontext
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

import openpyxl
import openpyxl.cell._writer as cell_writer
import pytest
from openpyxl.compat import safe_string
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, from_excel, to_excel

from excel_template_writer.xlsx import openpyxl_numeric_compat as compat


def _child(body: str) -> None:
    result = subprocess.run(
        [sys.executable, "-c", textwrap.dedent(body)],
        capture_output=True,
        text=True,
        timeout=30,
        env=os.environ.copy(),
    )
    assert result.returncode == 0, result.stdout + result.stderr


def _save(values: Sequence[object], *, epoch=CALENDAR_WINDOWS_1900, iso_dates=False) -> bytes:
    workbook = openpyxl.Workbook(iso_dates=iso_dates)
    workbook.epoch = epoch
    for index, value in enumerate(values, 1):
        workbook.active.cell(index, 1, value)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _tokens(data: bytes) -> list[str]:
    with ZipFile(BytesIO(data)) as package:
        root = ElementTree.fromstring(package.read("xl/worksheets/sheet1.xml"))
    nodes = root.findall(".//{*}v")
    assert all(node.text is not None for node in nodes)
    return [node.text for node in nodes if node.text is not None]


def test_stock_openpyxl_loses_regression_float_bits() -> None:
    _child("""
        from io import BytesIO
        from xml.etree import ElementTree
        from zipfile import ZipFile
        import openpyxl
        assert openpyxl.__version__ == '3.1.5'
        workbook = openpyxl.Workbook()
        workbook.active['A1'] = 100000.00000000001
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        with ZipFile(BytesIO(stream.getvalue())) as package:
            root = ElementTree.fromstring(package.read('xl/worksheets/sheet1.xml'))
        assert [node.text for node in root.findall('.//{*}v')] == ['100000']
    """)


def test_wrapper_round_trips_fixed_and_generated_binary64_values() -> None:
    formatter = compat._OwnedNumericFormatter(safe_string)
    values: list[float] = [
        100000.00000000001,
        1.0,
        0.0,
        -0.0,
        5e-324,
        sys.float_info.min,
        sys.float_info.max,
    ]
    generator = random.Random(3207)
    while len(values) < 1500:
        value = struct.unpack("!d", generator.getrandbits(64).to_bytes(8))[0]
        if math.isfinite(value):
            values.append(value)
    for value in values:
        assert type(value) is float
        token = formatter(value)
        assert token == repr(value)
        assert float(token).hex() == value.hex()


def test_wrapper_delegates_every_non_builtin_finite_float() -> None:
    class FloatSubclass(float):
        pass

    values = [
        12,
        True,
        False,
        Decimal("0.1"),
        None,
        "text",
        "=1+2",
        date(2026, 1, 1),
        datetime(2026, 1, 1),
        time(12),
        timedelta(days=1),
        float("inf"),
        float("-inf"),
        float("nan"),
        FloatSubclass(1.5),
    ]
    seen = []
    sentinel = object()

    def original(value):
        seen.append(value)
        return sentinel

    formatter = compat._OwnedNumericFormatter(original)
    for value in values:
        assert formatter(value) is sentinel
        assert seen[-1] is value
    assert len(seen) == len(values)


@pytest.mark.parametrize("writer", [cell_writer.etree_write_cell, cell_writer.lxml_write_cell])
def test_both_cell_writers_resolve_the_module_global_hook(monkeypatch, writer) -> None:
    class Recorder:
        def __init__(self):
            self.parts = []

        def element(self, *args):
            return nullcontext()

        def write(self, value):
            self.parts.append(value if isinstance(value, str) else ElementTree.tostring(value))

    monkeypatch.setattr(cell_writer, "safe_string", compat._OwnedNumericFormatter(safe_string))
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = 100000.00000000001
    recorder = Recorder()
    writer(recorder, workbook.active, workbook.active["A1"])
    assert "100000.00000000001" in "".join(
        part.decode() if isinstance(part, bytes) else part for part in recorder.parts
    )


@pytest.mark.parametrize("epoch", [CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904])
def test_temporal_serials_match_pinned_openpyxl_conversion(monkeypatch, epoch) -> None:
    values = [
        date(2026, 1, 1),
        time(12, 0, 0, 499),
        time(12, 0, 0, 500),
        time(12, 0, 0, 501),
        datetime(2026, 1, 1, 12, 0, 0, 499),
        datetime(2026, 1, 1, 12, 0, 0, 500),
        datetime(2026, 1, 1, 12, 0, 0, 501),
        timedelta(days=2, microseconds=500),
    ]
    monkeypatch.setattr(cell_writer, "safe_string", compat._OwnedNumericFormatter(safe_string))
    data = _save(values, epoch=epoch)
    assert _tokens(data) == [repr(to_excel(value, epoch)) for value in values]
    reopened = openpyxl.load_workbook(BytesIO(data))
    for index, value in enumerate(values, 1):
        expected = from_excel(to_excel(value, epoch), epoch, timedelta=isinstance(value, timedelta))
        assert reopened.active.cell(index, 1).value == expected
    reopened.close()

    calls = []

    def tracked(value):
        calls.append(value)
        return safe_string(value)

    monkeypatch.setattr(cell_writer, "safe_string", compat._OwnedNumericFormatter(tracked))
    iso_values = [date(2026, 1, 1), datetime(2026, 1, 1, 12, 0, 0, 123000), time(12, 0, 0, 123000)]
    iso_data = _save(iso_values, epoch=epoch, iso_dates=True)
    assert calls and all(type(value) is str for value in calls)
    assert _tokens(iso_data) == ["2026-01-01", "2026-01-01T12:00:00.123", "12:00:00.123"]
    reopened = openpyxl.load_workbook(BytesIO(iso_data))
    assert [reopened.active.cell(index, 1).value for index in range(1, 4)] == iso_values
    reopened.close()


def test_activation_idempotence_reload_and_raw_workbook_round_trips() -> None:
    _child("""
        import importlib
        from io import BytesIO
        import openpyxl
        import openpyxl.cell._writer as writer
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        assert compat.ensure_openpyxl_numeric_compatibility() is None
        wrapper = writer.safe_string
        assert compat.ensure_openpyxl_numeric_compatibility() is None
        importlib.reload(compat)
        assert compat.ensure_openpyxl_numeric_compatibility() is None
        assert writer.safe_string is wrapper
        for write_only in (False, True):
            book = openpyxl.Workbook(write_only=write_only)
            sheet = book.create_sheet() if write_only else book.active
            values = [100000.00000000001, 1.0, -0.0, 5e-324]
            for value in values:
                sheet.append([value])
            stream = BytesIO()
            book.save(stream)
            result = openpyxl.load_workbook(BytesIO(stream.getvalue()))
            for index, value in enumerate(values, 1):
                actual = result.active.cell(index, 1).value
                assert type(actual) is float and actual.hex() == value.hex()
            result.close()
    """)


@pytest.mark.parametrize(
    "mutation",
    [
        "openpyxl.__version__ = '3.1.6'",
        "del writer.safe_string",
        "writer.safe_string = lambda value: str(value)",
        "writer.etree_write_cell = lambda *args: None",
        "writer.safe_string = compat._OwnedNumericFormatter(lambda value: str(value))",
    ],
)
def test_incompatible_version_and_hooks_fail_closed(mutation) -> None:
    _child(f"""
        import openpyxl
        import openpyxl.cell._writer as writer
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        {mutation}
        before = getattr(writer, 'safe_string', None)
        failure = compat.ensure_openpyxl_numeric_compatibility()
        assert failure.code == {'"E3206"' if "__version__" in mutation else '"E3207"'}
        assert str(failure.location) == '<workbook>!A1'
        assert compat.ensure_openpyxl_numeric_compatibility() is failure
        assert getattr(writer, 'safe_string', None) is before
    """)


@pytest.mark.parametrize("replace_hook", [False, True])
def test_failed_self_test_restores_only_an_owned_hook(replace_hook) -> None:
    _child(f"""
        import openpyxl.cell._writer as writer
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        original = writer.safe_string
        foreign = lambda value: str(value)
        def fail():
            if {replace_hook!r}:
                writer.safe_string = foreign
            raise RuntimeError('deliberate self-test failure')
        compat._run_self_test = fail
        failure = compat.ensure_openpyxl_numeric_compatibility()
        assert failure.code == 'E3207'
        assert 'self-test' in failure.message
        assert writer.safe_string is (foreign if {replace_hook!r} else original)
        assert compat.ensure_openpyxl_numeric_compatibility() is failure
    """)


def test_post_activation_replacement_is_sticky_without_overwriting_foreign_hook() -> None:
    _child("""
        import openpyxl.cell._writer as writer
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        assert compat.ensure_openpyxl_numeric_compatibility() is None
        owned = writer.safe_string
        foreign = lambda value: str(value)
        writer.safe_string = foreign
        failure = compat.ensure_openpyxl_numeric_compatibility()
        assert failure.code == 'E3207'
        assert writer.safe_string is foreign
        writer.safe_string = owned
        assert compat.ensure_openpyxl_numeric_compatibility() is failure
    """)


def test_cached_failure_survives_module_reload() -> None:
    _child("""
        import importlib
        import openpyxl
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        openpyxl.__version__ = 'unverified'
        failure = compat.ensure_openpyxl_numeric_compatibility()
        assert failure.code == 'E3206'
        openpyxl.__version__ = '3.1.5'
        importlib.reload(compat)
        assert compat.ensure_openpyxl_numeric_compatibility() is failure
    """)


def test_fresh_module_recognizes_an_owned_hook_from_previous_incarnation() -> None:
    _child("""
        import importlib
        import sys
        import openpyxl.cell._writer as writer
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        assert compat.ensure_openpyxl_numeric_compatibility() is None
        wrapper = writer.safe_string
        name = compat.__name__
        del sys.modules[name]
        fresh = importlib.import_module(name)
        assert fresh is not compat
        assert fresh.ensure_openpyxl_numeric_compatibility() is None
        assert writer.safe_string is wrapper
    """)


@pytest.mark.parametrize("failure_kind", ["version", "self_test"])
def test_cached_failure_survives_a_fresh_module_incarnation(failure_kind) -> None:
    _child(f"""
        import importlib
        import sys
        import openpyxl
        import openpyxl.cell._writer as writer
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        original = writer.safe_string
        if {failure_kind!r} == 'version':
            openpyxl.__version__ = 'unverified'
        else:
            def fail():
                raise RuntimeError('deliberate self-test failure')
            compat._run_self_test = fail
        failure = compat.ensure_openpyxl_numeric_compatibility()
        assert failure.code == {'"E3206"' if failure_kind == "version" else '"E3207"'}
        openpyxl.__version__ = '3.1.5'
        name = compat.__name__
        del sys.modules[name]
        fresh = importlib.import_module(name)
        assert fresh is not compat
        assert fresh.ensure_openpyxl_numeric_compatibility() is failure
        assert compat.ensure_openpyxl_numeric_compatibility() is failure
        assert writer.safe_string is original
    """)


def test_module_incarnations_share_one_activation_self_test() -> None:
    _child("""
        import importlib
        import sys
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier, Event
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        name = compat.__name__
        del sys.modules[name]
        fresh = importlib.import_module(name)
        assert fresh is not compat
        start = Barrier(3)
        entered, release = Event(), Event()
        calls = []
        original = compat._run_self_test
        def gated():
            calls.append(1)
            entered.set()
            assert release.wait(10)
            original()
        compat._run_self_test = fresh._run_self_test = gated
        def activate(module):
            start.wait(10)
            return module.ensure_openpyxl_numeric_compatibility()
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(activate, module) for module in (compat, fresh)]
            start.wait(10)
            assert entered.wait(10)
            assert not any(future.done() for future in futures)
            release.set()
            assert [future.result(10) for future in futures] == [None, None]
        assert calls == [1]
    """)


@pytest.mark.parametrize("owner_kind", ["foreign", "malformed"])
def test_invalid_process_owner_fails_closed_without_replacement(owner_kind) -> None:
    _child(f"""
        import openpyxl.cell._writer as writer
        if {owner_kind!r} == 'malformed':
            from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
            foreign = writer._excel_template_writer_numeric_compat_state
            del foreign.lock
        else:
            foreign = object()
            writer._excel_template_writer_numeric_compat_state = foreign
        original = writer.safe_string
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        failure = compat.ensure_openpyxl_numeric_compatibility()
        assert failure.code == 'E3207'
        assert str(failure.location) == '<workbook>!A1'
        assert compat.ensure_openpyxl_numeric_compatibility() is failure
        assert writer._excel_template_writer_numeric_compat_state is foreign
        assert writer.safe_string is original
    """)


def test_real_self_test_rejects_a_broken_reader_and_restores_stock() -> None:
    _child("""
        import openpyxl
        import openpyxl.cell._writer as writer
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        original = writer.safe_string
        reader = openpyxl.load_workbook
        def broken_reader(*args, **kwargs):
            book = reader(*args, **kwargs)
            book.active['A1'] = 100000
            return book
        openpyxl.load_workbook = broken_reader
        failure = compat.ensure_openpyxl_numeric_compatibility()
        assert failure.code == 'E3207'
        assert 'self-test' in failure.message and 'float bits' in failure.message
        assert writer.safe_string is original
        assert compat.ensure_openpyxl_numeric_compatibility() is failure
    """)


def test_concurrent_activation_runs_one_self_test() -> None:
    _child("""
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier, Event
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        start = Barrier(9)
        entered, release = Event(), Event()
        calls = []
        original = compat._run_self_test
        def gated():
            calls.append(1)
            entered.set()
            assert release.wait(10)
            original()
        compat._run_self_test = gated
        def activate():
            start.wait(10)
            return compat.ensure_openpyxl_numeric_compatibility()
        with ThreadPoolExecutor(max_workers=8) as pool:
            futures = [pool.submit(activate) for _ in range(8)]
            start.wait(10)
            assert entered.wait(10)
            assert not any(future.done() for future in futures)
            release.set()
            assert [future.result(10) for future in futures] == [None] * 8
        assert calls == [1]
    """)


def test_activated_process_saves_independent_workbooks_concurrently() -> None:
    _child("""
        from concurrent.futures import ThreadPoolExecutor
        from io import BytesIO
        from threading import Barrier
        import openpyxl
        import openpyxl.writer.excel as excel_writer
        from excel_template_writer.xlsx import openpyxl_numeric_compat as compat
        assert compat.ensure_openpyxl_numeric_compatibility() is None
        overlap = Barrier(2)
        original = excel_writer.ExcelWriter.save
        def gated(self):
            overlap.wait(10)
            return original(self)
        excel_writer.ExcelWriter.save = gated
        def save():
            assert compat.ensure_openpyxl_numeric_compatibility() is None
            book = openpyxl.Workbook()
            book.active['A1'] = 100000.00000000001
            stream = BytesIO()
            book.save(stream)
            result = openpyxl.load_workbook(BytesIO(stream.getvalue()))
            assert result.active['A1'].value.hex() == (100000.00000000001).hex()
            result.close()
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(save) for _ in range(2)]
            for future in futures:
                future.result(15)
    """)
