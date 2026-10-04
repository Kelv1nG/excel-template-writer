from __future__ import annotations

import hashlib
import json
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from excel_template_writer.xlsx import render_workbook

REFERENCE_DIR = Path(__file__).parents[1] / "fixtures" / "numeric_fidelity"
REFERENCE = REFERENCE_DIR / "excel_numeric_reference.xlsx"
MANIFEST = REFERENCE_DIR / "excel_numeric_reference.json"
PRECISION_AMOUNTS = (
    "0.3",
    "0.31",
    "0.314",
    "0.3141",
    "0.31415",
    "0.314159",
    "0.3141592",
    "0.31415926",
    "0.314159265",
    "0.3141592653",
    "0.31415926535",
    "0.314159265358",
    "0.3141592653589",
    "0.31415926535897",
    "0.314159265358979",
)
OPERATIONS = (
    ("exact_add", "=0.5+0.25", "0.75"),
    ("exact_subtract", "=1.5-0.25", "1.25"),
    ("exact_multiply", "=1.125*2", "2.25"),
    ("exact_divide", "=5/2", "2.5"),
)


def _manifest() -> dict:
    assert MANIFEST.is_file(), "Frozen desktop-Excel manifest is missing"
    assert REFERENCE.is_file(), "Frozen desktop-Excel workbook is missing"
    return json.loads(MANIFEST.read_text(encoding="utf-8-sig"))


def _expected_cells() -> list[tuple[str, str, str, str]]:
    return [
        (f"precision_{digits:02}", f"B{digits + 1}", f"={amount}", amount)
        for digits, amount in enumerate(PRECISION_AMOUNTS, 1)
    ] + [
        (name, f"B{index}", formula, amount)
        for index, (name, formula, amount) in enumerate(OPERATIONS, 19)
    ]


def test_excel_reference_fixture_matches_recorded_provenance_and_checksum() -> None:
    manifest = _manifest()
    assert manifest["schema_version"] == 1
    assert manifest["producer"] == "Microsoft Excel desktop via PowerShell COM"
    assert manifest["workbook"] == REFERENCE.name
    assert manifest["sha256"] == hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
    for key in ("version", "build", "file_version", "executable"):
        assert manifest["excel"][key]
    assert manifest["locale"]["culture"]
    assert isinstance(manifest["locale"]["country_setting"], int)
    assert manifest["locale"]["decimal_separator"] in (".", ",")
    assert manifest["calculation"] == {
        "mode": "automatic",
        "mode_value": -4105,
        "precision_as_displayed": False,
        "iteration": False,
        "full_rebuild": True,
        "state_after_rebuild": 0,
        "automation_security": 3,
        "save_format": 51,
    }
    created = datetime.fromisoformat(manifest["created_utc"].replace("Z", "+00:00"))
    assert created.utcoffset() == timedelta(0)
    assert manifest["created_utc"].endswith("Z")

    workbook = load_workbook(REFERENCE, data_only=False)
    try:
        assert workbook.sheetnames == ["NumericReference"]
        assert workbook.calculation is not None
        # Excel omits these OOXML attributes when their default applies.
        assert workbook.calculation.fullPrecision in (None, True)
        assert workbook.calculation.iterate in (None, False)
        assert workbook.calculation.calcMode in (None, "auto")
        with ZipFile(REFERENCE) as archive:
            assert not any("vbaProject" in name for name in archive.namelist())
            properties = ElementTree.fromstring(archive.read("docProps/app.xml"))
            application = properties.find("{*}Application")
            assert application is not None and application.text == "Microsoft Excel"
    finally:
        workbook.close()


def test_excel_reference_contains_documented_formulas_and_cached_values() -> None:
    manifest = _manifest()
    assert len(manifest["cells"]) == 19
    formulas = load_workbook(REFERENCE, data_only=False)
    cached = load_workbook(REFERENCE, data_only=True)
    try:
        for record, (name, coordinate, formula, amount) in zip(
            manifest["cells"], _expected_cells(), strict=True
        ):
            assert (record["name"], record["cell"], record["formula"], record["decimal"]) == (
                name,
                coordinate,
                formula,
                amount,
            )
            assert record["sheet"] == "NumericReference"
            assert record["comparison_mode"] == (
                "decimal_amount" if name.startswith("precision_") else "binary64_exact"
            )
            assert list(formulas.defined_names[name].destinations) == [
                ("NumericReference", f"$B${coordinate[1:]}")
            ]
            cell = formulas["NumericReference"][coordinate]
            assert cell.data_type == "f" and cell.value == formula
            result = cached["NumericReference"][coordinate]
            assert result.data_type == "n" and type(result.value) is float
            assert Decimal(str(result.value)) == Decimal(amount)
            assert float(record["cached_decimal"]).hex() == result.value.hex()
            if record["comparison_mode"] == "binary64_exact":
                assert result.value.hex() == float(amount).hex()
    finally:
        formulas.close()
        cached.close()


def test_excel_reference_selected_xml_nodes_have_formula_and_cache() -> None:
    _manifest()
    with ZipFile(REFERENCE) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    cells = {cell.attrib["r"]: cell for cell in root.findall(".//{*}c")}
    for _, coordinate, formula, amount in _expected_cells():
        cell = cells[coordinate]
        assert cell.get("t") in (None, "n")
        formula_node = cell.find("{*}f")
        cache = cell.find("{*}v")
        assert formula_node is not None and formula_node.text == formula[1:]
        assert cache is not None and cache.text is not None
        # Excel may emit a longer binary64 token than the literal decimal amount.
        assert Decimal(str(float(cache.text))) == Decimal(amount)
        assert float(cache.text).hex() == float(amount).hex()


def test_rendered_decimal_precision_matrix_matches_frozen_excel_results(tmp_path: Path) -> None:
    _manifest()
    values = [Decimal(amount) for amount in PRECISION_AMOUNTS]
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = "Results"
        for index in range(15):
            sheet.cell(
                index + 1, 1, "{{ values[" + str(index) + "] }}"
            ).number_format = "0.000000000000000"
        workbook.save(template)
    finally:
        workbook.close()

    render_workbook(template, output, {"values": values})

    reference = load_workbook(REFERENCE, data_only=True)
    rendered = load_workbook(output, data_only=True)
    try:
        with ZipFile(output) as archive:
            root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        cells = {cell.attrib["r"]: cell for cell in root.findall(".//{*}c")}
        for index, amount in enumerate(values, 1):
            actual = rendered["Results"][f"A{index}"]
            expected = reference["NumericReference"][f"B{index + 1}"].value
            assert actual.data_type == "n" and type(actual.value) is float
            assert Decimal(str(actual.value)) == Decimal(str(expected)) == amount
            assert actual.number_format == "0.000000000000000"
            cell = cells[f"A{index}"]
            assert cell.get("t") == "n"
            token = cell.find("{*}v")
            assert token is not None and token.text == repr(actual.value)
            assert Decimal(token.text) == amount
    finally:
        reference.close()
        rendered.close()
