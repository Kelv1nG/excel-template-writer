from __future__ import annotations

from copy import copy
from datetime import date
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Protection, Side
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils.indexed_list import IndexedList

from excel_template_writer.diagnostics import (
    DiagnosticCode,
    SourceLocation,
    TemplateCompilationError,
)
from excel_template_writer.xlsx import render_workbook

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


class _ExplicitStyleCell(Cell):
    """Build an XLSX fixture with an explicit all-zero style, including blanks.

    Excel can write these records, but openpyxl normally omits them because its
    has_style tests nonzero component IDs rather than the workbook's default.
    This fixture-only override exercises the reader without patching XLSX XML.
    """

    @property
    def has_style(self) -> bool:
        return self._style is not None


def _set_default(workbook: Workbook, *, rich: bool = False) -> None:
    donor = Cell(workbook.active, row=1, column=1)
    donor.fill = PatternFill("solid", fgColor="FFFFFFFF")
    if rich:
        donor.font = Font(name="Arial", size=14, color="FF123456")
        donor.border = Border(bottom=Side(style="thin", color="FF112233"))
        donor.alignment = Alignment(horizontal="center", wrap_text=True)
        donor.number_format = '0.0000 "units"'
        donor.protection = Protection(locked=False, hidden=True)
        donor.quotePrefix = True
    workbook._cell_styles = IndexedList([copy(donor._style)])
    # Deliberately leave named Normal unchanged: it is not cellXfs[0].


def _default_cell(workbook: Workbook) -> Cell:
    donor = Cell(workbook.active, row=1, column=1)
    donor._style = copy(workbook._cell_styles[0])
    return donor


def _sheet_xml(path: Path) -> ElementTree.Element:
    with ZipFile(path) as archive:
        return ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))


def _style_fill(path: Path, *, normal: bool) -> tuple[str | None, str | None]:
    with ZipFile(path) as archive:
        styles = ElementTree.fromstring(archive.read("xl/styles.xml"))
    if normal:
        named_styles = styles.find("s:cellStyles", NS)
        style_xfs = styles.find("s:cellStyleXfs", NS)
        assert named_styles is not None and style_xfs is not None
        normal_ref = next(style for style in named_styles if style.get("builtinId") == "0")
        xf_id = normal_ref.get("xfId")
        assert xf_id is not None
        xf = style_xfs[int(xf_id)]
    else:
        cell_xfs = styles.find("s:cellXfs", NS)
        assert cell_xfs is not None
        xf = cell_xfs[0]
    fills = styles.find("s:fills", NS)
    assert fills is not None
    fill = fills[int(xf.get("fillId", "0"))]
    pattern = fill.find("s:patternFill", NS)
    assert pattern is not None
    color = pattern.find("s:fgColor", NS)
    return pattern.get("patternType"), None if color is None else color.get("rgb")


@pytest.mark.parametrize("default_white", [True, False])
def test_preserves_normal_style_independently_of_default_cell_xf(
    tmp_path: Path, default_white: bool
) -> None:
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    if default_white:
        _set_default(workbook)
    workbook._named_styles[0].fill = PatternFill("solid", fgColor="FFFFFFFF")
    workbook.active.sheet_view.showGridLines = True
    workbook.active["A1"] = "{{ title }}"
    workbook.save(template)
    workbook.close()

    render_workbook(template, output, {"title": "Report"})

    expected_default = ("solid", "FFFFFFFF") if default_white else (None, None)
    for path in (template, output):
        assert _style_fill(path, normal=True) == ("solid", "FFFFFFFF")
        assert _style_fill(path, normal=False) == expected_default
        assert _sheet_xml(path).find(".//s:c[@r='B2']", NS) is None
        reopened = load_workbook(path)
        try:
            assert reopened.active.sheet_view.showGridLines is True
        finally:
            reopened.close()


def test_preserves_custom_normal_style_components(tmp_path: Path) -> None:
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    normal = workbook._named_styles[0]
    normal.font = Font(name="Arial", size=14, bold=True, color="FF123456")
    normal.fill = PatternFill("solid", fgColor="FFFFFFFF")
    normal.border = Border(bottom=Side(style="thin", color="FF112233"))
    normal.alignment = Alignment(horizontal="center", wrap_text=True)
    normal.number_format = '0.0000 "units"'
    normal.protection = Protection(locked=False, hidden=True)
    workbook.active["A1"] = "{{ title }}"
    workbook.save(template)
    workbook.close()

    render_workbook(template, output, {"title": "Report"})

    rendered = load_workbook(output)
    try:
        normal = rendered._named_styles[0]
        assert normal.font.name == "Arial" and normal.font.sz == 14
        assert normal.font.bold is True and normal.font.color.rgb == "FF123456"
        assert normal.fill.patternType == "solid" and normal.fill.fgColor.rgb == "FFFFFFFF"
        assert normal.border.bottom.style == "thin"
        assert normal.border.bottom.color.rgb == "FF112233"
        assert normal.alignment.horizontal == "center" and normal.alignment.wrap_text is True
        assert normal.number_format == '0.0000 "units"'
        assert normal.protection.locked is False and normal.protection.hidden is True
    finally:
        rendered.close()


def test_normal_style_does_not_change_per_sheet_gridline_settings(tmp_path: Path) -> None:
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    workbook._named_styles[0].fill = PatternFill("solid", fgColor="FFFFFFFF")
    workbook.active.title = "Report"
    workbook.active.sheet_view.showGridLines = True
    workbook.active["A1"] = "{{ title }}"
    notes = workbook.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes["A1"] = "Read me"
    workbook.save(template)
    workbook.close()

    render_workbook(template, output, {"title": "Report"})

    rendered = load_workbook(output)
    try:
        assert rendered["Report"].sheet_view.showGridLines is True
        assert rendered["Notes"].sheet_view.showGridLines is False
        assert _style_fill(output, normal=True) == ("solid", "FFFFFFFF")
    finally:
        rendered.close()


def test_preserves_named_normal_when_builtin_id_is_absent(tmp_path: Path) -> None:
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    normal = workbook._named_styles[0]
    normal.builtinId = None
    normal.fill = PatternFill("solid", fgColor="FFFFFFFF")
    workbook.active["A1"] = "{{ title }}"
    workbook.save(template)
    workbook.close()

    render_workbook(template, output, {"title": "Report"})

    rendered = load_workbook(output)
    try:
        assert rendered.active["A1"].value == "Report"
        assert rendered._named_styles[0].fill.patternType == "solid"
        assert rendered._named_styles[0].fill.fgColor.rgb == "FFFFFFFF"
    finally:
        rendered.close()


@pytest.mark.parametrize("case", ["missing", "ambiguous"])
def test_rejects_unidentifiable_normal_style_without_publishing(tmp_path: Path, case: str) -> None:
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    if case == "missing":
        workbook._named_styles[0].name = "Custom"
        workbook._named_styles[0].builtinId = None
    else:
        workbook.add_named_style(NamedStyle(name="Duplicate Normal", builtinId=0))
    workbook.active["A1"] = "{{ title }}"
    workbook.save(template)
    workbook.close()

    with pytest.raises(TemplateCompilationError) as error:
        render_workbook(template, output, {"title": "Report"})

    assert DiagnosticCode.XLSX_NORMAL_STYLE_UNIDENTIFIABLE in {
        item.code for item in error.value.diagnostics
    }
    location = error.value.diagnostics[0].location
    assert isinstance(location, SourceLocation)
    assert location.sheet == "<workbook>"
    assert location.cell == "A1"
    assert not output.exists()


@pytest.mark.parametrize("rich", [False, True])
@pytest.mark.parametrize("items", [[], ["one", "two"]])
def test_preserves_workbook_default_without_materializing_blank_background(
    tmp_path: Path, rich: bool, items: list[str]
) -> None:
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    _set_default(workbook, rich=rich)
    sheet = workbook.active
    for address, text in {
        "A1": "{{ title }}",
        "A3": "Second text",
        "A5": "Third text",
        "A8": "{% for item in items %}",
        "B8": "{{ item }}",
        "C8": "{% endfor %}",
    }.items():
        sheet[address] = text
        sheet[address]._style = copy(workbook._cell_styles[0])
    sheet.row_dimensions[2].height = 27
    sheet.column_dimensions["B"].width = 24
    for dimension in (sheet.row_dimensions[2], sheet.column_dimensions["B"]):
        dimension._style = copy(workbook._cell_styles[0])
    sheet["E3"].fill = PatternFill("solid", fgColor="FFABCDEF")
    sheet["E3"].border = Border(left=Side(style="double", color="FF445566"))
    workbook.save(template)
    workbook.close()
    original_bytes = template.read_bytes()

    render_workbook(template, output, {"title": "First text", "items": items})

    assert template.read_bytes() == original_bytes
    rendered = load_workbook(output)
    try:
        default = _default_cell(rendered)
        assert default.fill.patternType == "solid"
        assert default.fill.fgColor.rgb == "FFFFFFFF"
        if rich:
            assert default.font.name == "Arial"
            assert default.font.sz == 14
            assert default.font.color.rgb == "FF123456"
            assert default.border.bottom.style == "thin"
            assert default.border.bottom.color.rgb == "FF112233"
            assert default.alignment.horizontal == "center"
            assert default.alignment.wrap_text is True
            assert default.number_format == '0.0000 "units"'
            assert default.protection.locked is False
            assert default.protection.hidden is True
            assert default.quotePrefix is True
        target = rendered.active
        assert [target[f"A{row}"].value for row in (1, 3, 5)] == [
            "First text",
            "Second text",
            "Third text",
        ]
        assert [target.cell(8 + i, 2).value for i in range(len(items))] == items
        assert target["E3"].fill.fgColor.rgb == "FFABCDEF"
        assert target["E3"].border.left.style == "double"
        assert target.row_dimensions[2].height == 27
        assert target.column_dimensions["B"].width == 24
    finally:
        rendered.close()
    for path in (template, output):
        xml = _sheet_xml(path)
        assert _style_fill(path, normal=True) == (None, None)
        assert _style_fill(path, normal=False) == ("solid", "FFFFFFFF")
        assert not xml.findall(".//s:c[@r='B2']", NS)
        assert not xml.findall(".//s:c[@r='H5']", NS)
        assert all("s" not in row.attrib for row in xml.findall("s:sheetData/s:row", NS))
        assert all("style" not in col.attrib for col in xml.findall("s:cols/s:col", NS))


def test_preserves_explicit_no_fill_cells_and_dimensions_over_white_default(tmp_path: Path) -> None:
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    _set_default(workbook)
    workbook._named_styles[0].fill = PatternFill("solid", fgColor="FFFFFFFF")
    sheet = workbook.active
    sheet["A1"] = "{{ title }}"
    sheet["A1"]._style = copy(workbook._cell_styles[0])
    for row, column, value in ((1, 3, "No fill"), (1, 4, None), (4, 2, "{{ item }}")):
        cell = _ExplicitStyleCell(sheet, row=row, column=column, value=value)
        cell._style = StyleArray()
        sheet._cells[(row, column)] = cell
    sheet["A4"] = "{% for item in items %}"
    sheet["C4"] = "{% endfor %}"
    sheet.row_dimensions[2].height = 22
    sheet.row_dimensions[2]._style = StyleArray()
    sheet.column_dimensions["E"].width = 20
    sheet.column_dimensions["E"]._style = StyleArray()
    workbook.save(template)
    workbook.close()
    source_xml = _sheet_xml(template)
    assert source_xml.find(".//s:c[@r='D1']", NS) is not None

    render_workbook(template, output, {"title": "White", "items": ["one", "two"]})

    rendered = load_workbook(output)
    try:
        assert _default_cell(rendered).fill.fgColor.rgb == "FFFFFFFF"
        target = rendered.active
        assert target["A1"].fill.fgColor.rgb == "FFFFFFFF"
        assert (1, 4) in target._cells  # Blank override must not disappear.
        for address in ("C1", "D1", "B4", "B5"):
            assert target[address].fill.patternType is None
            assert target[address].number_format == "General"
        assert target["C1"].value == "No fill"
        assert target["B4"].value == "one"
        assert target["B5"].value == "two"
        assert target.row_dimensions[2].fill.patternType is None
        assert target.column_dimensions["E"].fill.patternType is None
        assert target.row_dimensions[2].height == 22
        assert target.column_dimensions["E"].width == 20
    finally:
        rendered.close()
    xml = _sheet_xml(output)
    blank = xml.find(".//s:c[@r='D1']", NS)
    row = xml.find(".//s:row[@r='2']", NS)
    column = xml.find("s:cols/s:col", NS)
    assert blank is not None and blank.get("s") is not None
    assert row is not None and row.get("s") is not None
    assert column is not None and column.get("style") is not None


def test_fixed_unmaterialized_destinations_use_restored_default_after_value_assignment(
    tmp_path: Path,
) -> None:
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    _set_default(workbook, rich=True)
    sheet = workbook.active
    sheet["A1"] = '{% for item in items shift="none" %}{{ item }}{% endfor %}'
    sheet["A1"].fill = PatternFill("solid", fgColor="FFCCEEFF")
    workbook.save(template)
    workbook.close()

    render_workbook(template, output, {"items": ["one", date(2026, 9, 27)]})

    rendered = load_workbook(output)
    try:
        target = rendered.active
        assert target["A1"].fill.fgColor.rgb == "FFCCEEFF"
        assert target["A2"].fill.fgColor.rgb == "FFFFFFFF"
        assert target["A2"].number_format == '0.0000 "units"'
        assert target["A2"].font.name == "Arial"
        assert target["A2"].protection.locked is False
    finally:
        rendered.close()
