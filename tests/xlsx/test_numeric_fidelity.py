from __future__ import annotations

import math
import struct
import sys
from collections.abc import Sequence
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from openpyxl import Workbook, load_workbook

from excel_template_writer.compiler import compile_sheet
from excel_template_writer.diagnostics import (
    Diagnostic,
    DiagnosticCode,
    SourceLocation,
    TemplateRenderError,
)
from excel_template_writer.limits import DEFAULT_RESOURCE_LIMITS
from excel_template_writer.render import render_sheet
from excel_template_writer.xlsx import render_workbook, writer
from excel_template_writer.xlsx.reader import read_workbook
from excel_template_writer.xlsx.validation import plan_sheet_features


def _template(path: Path, count: int, formats: Sequence[str] = ()) -> Path:
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = "Numbers"
        for index in range(count):
            cell = sheet.cell(index + 1, 1, "{{ values[" + str(index) + "] }}")
            if formats:
                cell.number_format = formats[index]
        workbook.save(path)
    finally:
        workbook.close()
    return path


def _xml_cells(path: Path) -> dict[str, ElementTree.Element]:
    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    return {cell.attrib["r"]: cell for cell in root.findall(".//{*}c")}


def _assert_numeric_output(path: Path, values: Sequence[float]) -> None:
    cells = _xml_cells(path)
    reopened = load_workbook(path)
    try:
        for index, value in enumerate(values, 1):
            assert type(value) is float
            coordinate = f"A{index}"
            xml_cell = cells[coordinate]
            assert xml_cell.get("t") == "n"
            token = xml_cell.find("{*}v")
            assert token is not None and token.text == repr(value)
            cell = reopened["Numbers"][coordinate]
            assert cell.data_type == "n"
            assert type(cell.value) is float
            assert cell.value.hex() == value.hex(), coordinate
    finally:
        reopened.close()


def test_render_workbook_round_trips_fixed_float_matrix_bit_exactly(tmp_path: Path) -> None:
    positive = [
        1.0,
        0.0,
        100000.00000000001,
        0.30000000000000004,
        1.2345678901234567,
        1e20,
        1e-20,
        math.nextafter(1.0, 0.0),
        math.nextafter(1.0, math.inf),
        math.nextafter(0.0, math.inf),
        math.nextafter(sys.float_info.min, 0.0),
        sys.float_info.min,
        sys.float_info.max,
    ]
    values = positive + [-value for value in positive]
    template = _template(tmp_path / "template.xlsx", len(values))
    output = tmp_path / "output.xlsx"

    result = render_workbook(template, output, {"values": values})

    assert result.output_path == output
    assert result.diagnostics == ()
    _assert_numeric_output(output, values)


def test_float_xml_tokens_use_repr_and_remain_numeric(tmp_path: Path) -> None:
    values = [1.0, 0.0, -0.0, 100000.00000000001, 0.30000000000000004, 5e-324]
    template = _template(tmp_path / "template.xlsx", len(values))
    output = tmp_path / "output.xlsx"

    render_workbook(template, output, {"values": values})

    _assert_numeric_output(output, values)


_finite_raw_bits = st.binary(min_size=8, max_size=8).filter(
    lambda bits: (int.from_bytes(bits) >> 52) & 0x7FF != 0x7FF
)


@settings(
    max_examples=30,
    deadline=None,
    suppress_health_check=[HealthCheck.function_scoped_fixture],
)
@given(patterns=st.lists(_finite_raw_bits, min_size=1, max_size=24))
def test_render_workbook_round_trips_generated_finite_float_batches(
    tmp_path: Path, patterns: list[bytes]
) -> None:
    values = [struct.unpack("!d", bits)[0] for bits in patterns]
    template = _template(tmp_path / "template.xlsx", len(values))
    output = tmp_path / "output.xlsx"

    render_workbook(template, output, {"values": values})

    _assert_numeric_output(output, values)


def test_supported_decimal_is_numeric_in_xml_and_preserves_number_format(tmp_path: Path) -> None:
    values = [
        Decimal((sign, (*map(int, "314159265358979"[:digits]), 0, 0), exponent - 2))
        for digits in range(1, 16)
        for sign in (0, 1)
        for exponent in (-20, -3, 0, 20)
    ] + [
        Decimal("0"),
        Decimal("-0E-1000"),
        Decimal("2.22507385850721E-308"),
        Decimal("-2.22507385850721E-308"),
        Decimal("1.79769313486231E+308"),
        Decimal("-1.79769313486231E+308"),
    ]
    formats = ["0.0000" if index % 2 else "#,##0.00;[Red]-#,##0.00" for index in range(len(values))]
    template = _template(tmp_path / "template.xlsx", len(values), formats)
    output = tmp_path / "output.xlsx"

    render_workbook(template, output, {"values": values})

    _assert_numeric_output(output, [float(value) for value in values])
    cells = _xml_cells(output)
    reopened = load_workbook(output)
    try:
        for index, (value, number_format) in enumerate(zip(values, formats, strict=True), 1):
            token = cells[f"A{index}"].find("{*}v")
            assert token is not None and token.text is not None
            assert Decimal(token.text) == value
            assert reopened["Numbers"][f"A{index}"].number_format == number_format
    finally:
        reopened.close()


def test_decimal_scale_variants_have_equal_amounts_and_authored_formats(tmp_path: Path) -> None:
    values = [Decimal(text) for text in ("12.5", "12.50", "12.5000", "1.25E+1")]
    formats = ["0.0", "0.00", "0.0000", "0.00E+00"]
    template = _template(tmp_path / "template.xlsx", len(values), formats)
    output = tmp_path / "output.xlsx"

    render_workbook(template, output, {"values": values})

    _assert_numeric_output(output, [12.5] * 4)
    reopened = load_workbook(output)
    try:
        assert [cell.number_format for cell in reopened["Numbers"]["A"]] == formats
    finally:
        reopened.close()


def test_decimal_policy_does_not_reject_unused_or_mixed_text_values(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Amount: {{ amount }}"
    workbook.save(template)
    workbook.close()
    output = tmp_path / "output.xlsx"

    render_workbook(
        template,
        output,
        {"amount": Decimal("1234567890123456"), "unused": Decimal("1E+1000")},
    )

    reopened = load_workbook(output)
    try:
        assert reopened.active["A1"].value == "Amount: 1234567890123456"
        assert reopened.active["A1"].data_type == "s"
    finally:
        reopened.close()


@pytest.mark.parametrize(
    ("value", "code"),
    [
        (Decimal("1234567890123456"), DiagnosticCode.XLSX_DECIMAL_PRECISION_EXCEEDED),
        (Decimal("1E+309"), DiagnosticCode.XLSX_DECIMAL_OUT_OF_RANGE),
        (Decimal("1E-324"), DiagnosticCode.XLSX_DECIMAL_OUT_OF_RANGE),
        (Decimal("1E-308"), DiagnosticCode.XLSX_DECIMAL_OUT_OF_RANGE),
    ],
)
def test_decimal_preflight_runs_before_output_directory_creation(
    tmp_path: Path, value: Decimal, code: DiagnosticCode
) -> None:
    template = _template(tmp_path / "template.xlsx", 1)
    output = tmp_path / "absent" / "nested" / "output.xlsx"

    with pytest.raises(TemplateRenderError) as caught:
        render_workbook(template, output, {"values": [value]})

    assert [(item.code, item.location) for item in caught.value.diagnostics] == [
        (code, SourceLocation("Numbers", "A1"))
    ]
    assert not output.parent.parent.exists()


def test_decimal_preflight_preserves_existing_output_bytes(tmp_path: Path) -> None:
    template = _template(tmp_path / "template.xlsx", 2)
    output = tmp_path / "output.xlsx"
    original = b"previous published output"
    output.write_bytes(original)

    with pytest.raises(TemplateRenderError) as caught:
        render_workbook(
            template, output, {"values": [Decimal("1234567890123456"), Decimal("1E-324")]}
        )

    assert [(item.code, item.location) for item in caught.value.diagnostics] == [
        (DiagnosticCode.XLSX_DECIMAL_PRECISION_EXCEEDED, SourceLocation("Numbers", "A1")),
        (DiagnosticCode.XLSX_DECIMAL_OUT_OF_RANGE, SourceLocation("Numbers", "A2")),
    ]
    assert output.read_bytes() == original
    assert set(tmp_path.iterdir()) == {template, output}


def test_repeated_invalid_decimals_report_each_destination_coordinate(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    for title in ("First", "Second"):
        sheet = workbook.active if title == "First" else workbook.create_sheet()
        sheet.title = title
        sheet["A1"] = "Header"
        sheet["A2"] = "{% for row in rows %}{{ row.precise }}"
        sheet["B2"] = "{{ row.tiny }}{% endfor %}"
    workbook.save(template)
    workbook.close()
    output = tmp_path / "output.xlsx"
    row = {"precise": Decimal("1234567890123456"), "tiny": Decimal("1E-324")}

    with pytest.raises(TemplateRenderError) as caught:
        render_workbook(template, output, {"rows": [row, row, row]})

    assert [(item.code, item.location) for item in caught.value.diagnostics] == [
        (code, SourceLocation(sheet, f"{column}{index}"))
        for sheet in ("First", "Second")
        for index in (2, 3, 4)
        for column, code in (
            ("A", DiagnosticCode.XLSX_DECIMAL_PRECISION_EXCEEDED),
            ("B", DiagnosticCode.XLSX_DECIMAL_OUT_OF_RANGE),
        )
    ]
    assert not output.exists()


def test_internal_writer_preflights_before_path_or_workbook_creation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    template = _template(tmp_path / "template.xlsx", 1)
    snapshot = read_workbook(template)
    sheet = snapshot.sheets[0]
    compiled = compile_sheet(sheet.template).require()
    plan = render_sheet(compiled, {"values": [Decimal("1234567890123456")]}).require()
    feature_plan, diagnostics = plan_sheet_features(
        sheet, compiled, plan, worksheet_names=frozenset({"Numbers"})
    )
    assert diagnostics == ()

    def forbidden(*args: object, **kwargs: object) -> None:
        pytest.fail("writer touched destination before numeric preflight")

    monkeypatch.setattr(writer, "Path", forbidden)
    monkeypatch.setattr(writer, "Workbook", forbidden)
    with pytest.raises(TemplateRenderError) as caught:
        writer.write_workbook(
            snapshot, (plan,), (feature_plan,), "output.xlsx", limits=DEFAULT_RESOURCE_LIMITS
        )

    assert caught.value.diagnostics[0].code == DiagnosticCode.XLSX_DECIMAL_PRECISION_EXCEEDED
    assert plan.cells[0].value == Decimal("1234567890123456")


@pytest.mark.parametrize("existing", (False, True))
def test_openpyxl_compatibility_failure_is_atomic_before_workbook_creation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, existing: bool
) -> None:
    template = _template(tmp_path / "template.xlsx", 1)
    output = tmp_path / "output.xlsx" if existing else tmp_path / "absent" / "output.xlsx"
    original = b"previous published output"
    if existing:
        output.write_bytes(original)
    diagnostic = Diagnostic(
        DiagnosticCode.XLSX_OPENPYXL_COMPAT_FAILED,
        "numeric formatter ownership check failed",
        SourceLocation("<workbook>", "A1"),
    )

    def forbidden(*args: object, **kwargs: object) -> None:
        pytest.fail("writer touched destination before compatibility validation")

    monkeypatch.setattr(
        writer, "ensure_openpyxl_numeric_compatibility", lambda: diagnostic, raising=False
    )
    monkeypatch.setattr(writer, "Path", forbidden)
    monkeypatch.setattr(writer, "Workbook", forbidden)
    with pytest.raises(TemplateRenderError) as caught:
        render_workbook(template, output, {"values": [1.0]})

    assert caught.value.diagnostics == (diagnostic,)
    if existing:
        assert output.read_bytes() == original
    else:
        assert not output.parent.exists()
